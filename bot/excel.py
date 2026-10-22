"""Excel (.xls / .xlsx) import va export."""
import io
from dataclasses import dataclass, field
from datetime import date, datetime

from . import db
from .greeting import age_on, parse_date
from .translit import search_key

NAME_HEADERS = ("сотрудник", "фио", "ф.и.о", "xodim", "f.i.sh", "fio", "ism")
DATE_HEADERS = ("дата рождения", "tug", "туғ", "birth")
GENDER_HEADERS = ("пол", "jins", "жинс", "gender")


@dataclass
class ParsedRow:
    full_name: str
    birth_date: date
    gender: str


@dataclass
class ParseResult:
    rows: list[ParsedRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def _read_cells(data: bytes, filename: str) -> list[list]:
    if filename.lower().endswith(".xls"):
        import xlrd

        book = xlrd.open_workbook(file_contents=data)
        sheet = book.sheet_by_index(0)
        rows = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                else:
                    row.append(cell.value)
            rows.append(row)
        return rows
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]


def _find_col(header: list, keys: tuple) -> int | None:
    for i, v in enumerate(header):
        s = str(v or "").strip().lower()
        if any(k in s for k in keys):
            return i
    return None


def _gender(value) -> str | None:
    s = str(value or "").strip().lower()
    if s.startswith(("муж", "erk", "эрк", "m", "м")):
        return "M"
    if s.startswith(("жен", "ayol", "аёл", "f", "ж", "a")):
        return "F"
    return None


def parse(data: bytes, filename: str) -> ParseResult:
    cells = _read_cells(data, filename)
    res = ParseResult()
    header_idx = None
    for i, row in enumerate(cells[:30]):
        if _find_col(row, NAME_HEADERS) is not None and _find_col(row, DATE_HEADERS) is not None:
            header_idx = i
            break
    if header_idx is None:
        res.errors.append("Sarlavha qatori topilmadi (F.I.Sh, Tug‘ilgan sana, Jins ustunlari kerak).")
        return res
    header = cells[header_idx]
    c_name, c_date, c_gender = (
        _find_col(header, NAME_HEADERS),
        _find_col(header, DATE_HEADERS),
        _find_col(header, GENDER_HEADERS),
    )
    if c_gender is None:
        res.errors.append("«Jins» (Пол) ustuni topilmadi.")
        return res
    for n, row in enumerate(cells[header_idx + 1:], start=header_idx + 2):
        if max(c_name, c_date, c_gender) >= len(row):
            continue
        name = str(row[c_name] or "").strip()
        if not name:
            continue
        raw_date = row[c_date]
        if isinstance(raw_date, datetime):
            bdate = raw_date.date()
        elif isinstance(raw_date, date):
            bdate = raw_date
        else:
            bdate = parse_date(str(raw_date or ""))
        gender = _gender(row[c_gender])
        if not bdate or not gender:
            res.errors.append(f"{n}-qator: {name} — sana yoki jins noto‘g‘ri")
            continue
        res.rows.append(ParsedRow(name, bdate, gender))
    return res


def _identity(name: str, bdate: str) -> tuple[str, str]:
    return search_key(name), bdate


def plan(rows: list[ParsedRow], sync: bool) -> dict:
    """Bazaga nima qo‘shilishi / o‘chirilishini hisoblaydi (hali yozmaydi)."""
    existing = {_identity(e["full_name"], e["birth_date"]): e for e in db.list_employees()}
    incoming = {}
    for r in rows:
        incoming.setdefault(_identity(r.full_name, r.birth_date.isoformat()), r)
    to_add = [r for k, r in incoming.items() if k not in existing]
    to_update = [
        (existing[k]["id"], r.gender)
        for k, r in incoming.items()
        if k in existing and existing[k]["gender"] != r.gender
    ]
    to_delete = [e for k, e in existing.items() if k not in incoming] if sync else []
    return {"add": to_add, "update": to_update, "delete": to_delete}


def apply(p: dict) -> None:
    for r in p["add"]:
        db.add_employee(r.full_name, r.birth_date, r.gender)
    for emp_id, gender in p["update"]:
        db.update_employee(emp_id, gender=gender)
    for e in p["delete"]:
        db.delete_employee(e["id"])


def export() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Xodimlar"
    ws.append(["№", "F.I.Sh", "Tug‘ilgan sana", "Yosh", "Jins", "Telegram"])
    for c in ws[1]:
        c.font = Font(bold=True)
    today = date.today()
    for i, e in enumerate(db.list_employees(), start=1):
        ws.append([
            i,
            e["full_name"],
            date.fromisoformat(e["birth_date"]).strftime("%d.%m.%Y"),
            age_on(e, today),
            "Erkak" if e["gender"] == "M" else "Ayol",
            e["telegram"],
        ])
    for col, width in zip("ABCDEF", (6, 42, 16, 7, 9, 20)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
